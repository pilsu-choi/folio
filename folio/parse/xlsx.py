"""XLSX native parsing (FR-28): sheet = page, data region = table. Formulas give the cached (stored) value.

Header rows per table: Excel table (ListObject) `headerRowCount` > sheet print titles > VLM text task
`table_header` (when enabled); otherwise `header_rows` is left unset.
"""

from __future__ import annotations

import datetime as dt
import io
import re

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from folio.parse import vlm as vlm_mod
from folio.parse.assemble import table_text


def _fmt(cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        fmt = cell.number_format or ""
        if "%" in fmt:
            dec = len(fmt.split(".")[1].rstrip("%")) if "." in fmt else 0
            return f"{v * 100:.{dec}f}%"
        return str(int(v)) if v == int(v) and abs(v) < 1e15 else f"{v:.12g}"
    return str(v).strip()


def _regions(ws) -> list[tuple[int, int, int, int]]:
    """Connected occupied areas (4-neighbourhood, merged ranges count as occupied) merged into bounding boxes."""
    occ: set[tuple[int, int]] = set()
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                occ.add((c.row, c.column))
    for m in ws.merged_cells.ranges:
        if (m.min_row, m.min_col) in occ:
            occ.update((r, c) for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1))
    seen: set[tuple[int, int]] = set()
    rects = []
    for start in sorted(occ):
        if start in seen:
            continue
        stack, comp = [start], []
        seen.add(start)
        while stack:
            r, c = stack.pop()
            comp.append((r, c))
            for nb in ((r + 1, c), (r - 1, c), (r, c + 1), (r, c - 1)):
                if nb in occ and nb not in seen:
                    seen.add(nb)
                    stack.append(nb)
        rects.append([min(r for r, _ in comp), min(c for _, c in comp), max(r for r, _ in comp), max(c for _, c in comp)])
    merged = True
    while merged:  # bounding boxes of different components may overlap (L shapes): union them
        merged = False
        for i in range(len(rects)):
            for j in range(i + 1, len(rects)):
                a, b = rects[i], rects[j]
                if a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]:
                    rects[i] = [min(a[0], b[0]), min(a[1], b[1]), max(a[2], b[2]), max(a[3], b[3])]
                    del rects[j]
                    merged = True
                    break
            if merged:
                break
    return sorted(tuple(r) for r in rects)


def _table(ws, rect) -> dict:
    r0, c0, r1, c1 = rect
    grid = [[_fmt(ws.cell(r, c)) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]
    spans = []
    for m in ws.merged_cells.ranges:
        if r0 <= m.min_row and m.max_row <= r1 and c0 <= m.min_col and m.max_col <= c1:
            spans.append({"row": m.min_row - r0, "col": m.min_col - c0, "rowspan": m.max_row - m.min_row + 1,
                          "colspan": m.max_col - m.min_col + 1})
    b = {"type": "table", "text": table_text(grid), "cells": grid, "text_source": "native"}
    if spans:
        b["spans"] = sorted(spans, key=lambda s: (s["row"], s["col"]))
    return b


HEADER_ASK_ROWS = 15  # rows sent to the model; a header deeper than this is not detected
_HEADER_RE = re.compile(r'"?header_rows"?\s*[:=]\s*(\d+)')


def _native_header(ws, rect) -> int | None:
    """Header rows declared in the file: top-most Excel table inside the region, else print title rows."""
    r0, c0, r1, c1 = rect
    tops = []
    for t in ws.tables.values():
        tc0, tr0, tc1, tr1 = range_boundaries(t.ref)
        if r0 <= tr0 and tr1 <= r1 and c0 <= tc0 and tc1 <= c1:
            tops.append((tr0, 1 if t.headerRowCount is None else t.headerRowCount))
    if tops:
        tr0, n = min(tops)
        return tr0 - r0 + n
    m = re.fullmatch(r"\$?(\d+):\$?(\d+)", ws.print_title_rows or "")
    if m and int(m[1]) <= r0 <= int(m[2]):
        return min(int(m[2]), r1) - r0 + 1
    return None


def _ask_header(grid: list[list[str]], sess, page: int) -> int | None:
    rows = grid[:HEADER_ASK_ROWS]
    text = "\n".join(f"{i}: " + " | ".join(c[:40] for c in r) for i, r in enumerate(rows))
    out = sess.ask("table_header", None, page, text=text)
    if out is None:  # failure already recorded by the session
        return None
    found = _HEADER_RE.findall(out)  # last match: reasoning text may precede the answer
    if not found:
        sess.note(page, "table_header", "no header_rows in response")
        return None
    k = int(found[-1])
    return k if k < len(rows) else None


def _chart_title(ch) -> str:
    try:
        return "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or [])).strip()
    except Exception:
        return ""


def _header_rows(ws, rect, grid, sess, ask: bool, page: int) -> int | None:
    k = _native_header(ws, rect)
    if k is None and ask and sess.text_enabled and len(grid) >= 2 and len(grid[0]) >= 2:
        k = _ask_header(grid, sess, page)
    return k


def parse_xlsx(data: bytes, include_hidden: bool, sess, header_llm: bool = True) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    pages = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and not include_hidden:
            continue
        n = len(pages) + 1
        sec = [ws.title]
        blocks = [{"type": "heading", "level": 1, "text": ws.title, "text_source": "native", "section_path": []}]
        for rect in _regions(ws):
            b = {**_table(ws, rect), "section_path": sec}
            k = _header_rows(ws, rect, b["cells"], sess, header_llm, n)
            if k is not None:
                b["header_rows"] = k
            blocks.append(b)
        figs = []
        for img in getattr(ws, "_images", []):
            try:
                figs.append(("", vlm_mod.to_png(img._data())))
            except Exception:
                figs.append(("", None))
        figs += [(_chart_title(ch), None) for ch in getattr(ws, "_charts", [])]
        for text, png in figs:
            src = "native"
            desc = sess.ask("figure", png, n)
            if desc:
                text, src = desc.strip(), "vlm"
            blocks.append({"type": "figure", "text_source": src, "section_path": sec, **({"text": text} if text else {})})
        pages.append({"page": n, "width": None, "height": None, "blocks": blocks, "quality": None})
    return pages
