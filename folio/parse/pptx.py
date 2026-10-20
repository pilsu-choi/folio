"""PPTX native parsing (FR-27): slide = page, shape positions -> bbox, reading order top->bottom, left->right."""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from folio.parse import vlm as vlm_mod
from folio.parse.assemble import apply_sections, table_text

EMU_PT = 12700
_TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE)
_BODY = (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)
_NOTES = "노트"


def convert_to_pdf(data: bytes) -> bytes:
    """`pptx_mode="convert"`: LibreOffice -> PDF (then the regular PDF path)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise NotImplementedError("pptx convert mode needs LibreOffice; not bundled")
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "in.pptx"
        src.write_bytes(data)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", d, str(src)],
                       check=True, capture_output=True, timeout=300)
        return (Path(d) / "in.pdf").read_bytes()


def _box(sh) -> list[float] | None:
    try:
        if sh.left is None or sh.top is None or sh.width is None or sh.height is None:
            return None
        return [sh.left / EMU_PT, sh.top / EMU_PT, (sh.left + sh.width) / EMU_PT, (sh.top + sh.height) / EMU_PT]
    except Exception:
        return None


def _ph_type(sh):
    try:
        return sh.placeholder_format.type if sh.is_placeholder else None
    except Exception:
        return None


def _flat(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _flat(sh.shapes)
        else:
            yield sh


def _is_bullet(p, ph) -> bool:
    ppr = p._p.pPr
    if ppr is not None:
        tags = {c.tag.rsplit("}", 1)[-1] for c in ppr}
        if "buNone" in tags:
            return False
        if tags & {"buChar", "buAutoNum", "buBlip"}:
            return True
    return p.level > 0 or ph in _BODY


def _text_blocks(sh, ph) -> list[dict]:
    box = _box(sh)
    out: list[dict] = []
    items: list[str] = []

    def emit(b: dict) -> None:
        b["text_source"] = "native"
        if box:
            b["bbox"] = list(box)
        out.append(b)

    def flush() -> None:
        if items:
            emit({"type": "list", "text": "\n".join(items)})
            items.clear()

    for p in sh.text_frame.paragraphs:
        text = "".join(r.text for r in p.runs).strip() if p.runs else p.text.strip()
        text = text.replace("\x0b", "\n")
        if not text:
            continue
        if _is_bullet(p, ph):
            items.append("  " * p.level + "- " + text)
        else:
            flush()
            emit({"type": "paragraph", "text": text})
    flush()
    return out


def _table_block(sh) -> dict:
    tbl = sh.table
    grid, spans = [], []
    for r, row in enumerate(tbl.rows):
        cells = []
        for c, cell in enumerate(row.cells):
            cells.append("" if cell.is_spanned else " ".join(cell.text.split()))
            if cell.is_merge_origin:
                spans.append({"row": r, "col": c, "rowspan": cell.span_height, "colspan": cell.span_width})
        grid.append(cells)
    b = {"type": "table", "text": table_text(grid), "cells": grid, "text_source": "native"}
    if spans:
        b["spans"] = spans
    return b


def _chart_text(chart) -> str:
    parts = []
    try:
        if chart.has_title:
            parts.append(chart.chart_title.text_frame.text.strip())
    except Exception:
        pass
    try:
        names = [s.name for plot in chart.plots for s in plot.series if s.name]
        if names:
            parts.append("Series: " + ", ".join(names))
    except Exception:
        pass
    return ". ".join(p for p in parts if p)


def _figure(sh, n: int, sess) -> dict | None:
    text, src, png = "", "native", None
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        descr = sh._element.xpath(".//p:cNvPr/@descr")
        text = (descr[0] if descr else "").strip()
        png = vlm_mod.to_png(sh.image.blob)
    elif getattr(sh, "has_chart", False) and sh.has_chart:
        text = _chart_text(sh.chart)
    else:
        return None
    desc = sess.ask("figure", png, n)
    if desc:
        text, src = desc.strip(), "vlm"
    b = {"type": "figure", "text_source": src}
    if text:
        b["text"] = text
    return b


def parse_pptx(data: bytes, notes: bool, sess) -> list[dict]:
    prs = Presentation(io.BytesIO(data))
    w, h = prs.slide_width / EMU_PT, prs.slide_height / EMU_PT
    pages = []
    for n, slide in enumerate(prs.slides, 1):
        title_sh = slide.shapes.title
        blocks: list[dict] = []
        if title_sh is not None and title_sh.has_text_frame and title_sh.text_frame.text.strip():
            blocks.append({"type": "heading", "level": 1, "text": " ".join(title_sh.text_frame.text.split()),
                           "text_source": "native", **({"bbox": _box(title_sh)} if _box(title_sh) else {})})
        title_el = title_sh.element if title_sh is not None else None
        shapes = sorted((s for s in _flat(slide.shapes) if s.element is not title_el),
                        key=lambda s: ((_box(s) or [0, 0])[1], (_box(s) or [0, 0])[0]))
        for sh in shapes:
            ph = _ph_type(sh)
            if ph in _TITLES and sh.has_text_frame and sh.text_frame.text.strip() and not blocks:
                blocks.append({"type": "heading", "level": 1, "text": " ".join(sh.text_frame.text.split()),
                               "text_source": "native", **({"bbox": _box(sh)} if _box(sh) else {})})
                continue
            new: list[dict] = []
            if getattr(sh, "has_table", False) and sh.has_table:
                new = [_table_block(sh)]
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "has_chart", False) and sh.has_chart:
                f = _figure(sh, n, sess)
                new = [f] if f else []
            elif sh.has_text_frame:
                new = _text_blocks(sh, ph)
            if _box(sh):
                for b in new:
                    b.setdefault("bbox", _box(sh))
            blocks.extend(new)
        stack = apply_sections(blocks)
        if notes and slide.has_notes_slide:
            for p in slide.notes_slide.notes_text_frame.text.split("\n"):
                if p.strip():
                    blocks.append({"type": "paragraph", "text": p.strip(), "text_source": "native",
                                   "section_path": [t for _, t in stack] + [_NOTES]})
        for b in blocks:
            if "bbox" in b:
                b["bbox"] = [round(x, 2) for x in b["bbox"]]
        pages.append({"page": n, "width": round(w, 2), "height": round(h, 2), "blocks": blocks, "quality": None})
    return pages
