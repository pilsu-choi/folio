"""VLM default for image content (FR-02/05/06): auto resolution, fallbacks, call counts. Stub server, no network."""

from __future__ import annotations

import httpx
import pytest
from fastapi.testclient import TestClient

from folio.backends import vlm_stub
from folio.backends.base import PageStructure, Region, Table, TableCell
from folio.parse import default
from tests._parse_helpers import FIX, FakeBackend, ocr_words_for, validate_doc


@pytest.fixture(autouse=True)
def _no_vlm_env(monkeypatch):
    monkeypatch.delenv("FOLIO_VLM_URL", raising=False)
    monkeypatch.delenv("FOLIO_VLM_MODEL", raising=False)


@pytest.fixture
def stub():
    seen = []
    app = vlm_stub.create_app()

    @app.middleware("http")
    async def spy(request, call_next):
        if request.method == "POST":
            seen.append(request.url.path)
        return await call_next(request)

    return {"url": "http://testserver/v1", "model": "stub", "client": TestClient(app)}, seen


def srcs(doc):
    return {b.text_source for _, b in doc.iter_blocks()}


def test_auto_without_endpoint_is_disabled_and_uses_ocr():
    fb = FakeBackend(ocr_words_for(["스캔 문서", "OCR text"]))
    doc = default(str(FIX / "mixed.pdf"), backend=fb)  # vlm="auto" default
    assert doc.meta["vlm"] == {"enabled": False, "reason": "FOLIO_VLM_URL not set"}
    assert doc.meta["calls"]["vlm"] == 0 and doc.meta["vlm_fallbacks"] == [] and doc.meta["errors"] == []
    assert [{b.text_source for b in p.blocks} for p in doc.pages] == [{"pdf_text"}, {"ocr"}, {"ocr"}]


def test_auto_with_env_endpoint_enabled(monkeypatch):
    monkeypatch.setenv("FOLIO_VLM_URL", "http://vlm.invalid/v1")
    monkeypatch.setenv("FOLIO_VLM_MODEL", "m")
    doc = default(str(FIX / "digital.pdf"), backend=FakeBackend())
    assert doc.meta["vlm"] == {"enabled": True, "model": "m"} and doc.meta["calls"]["vlm"] == 0


def test_scanned_pages_and_images_use_vlm(stub):
    cfg, seen = stub
    doc = default(str(FIX / "mixed.pdf"), backend=FakeBackend(), vlm=cfg, vlm_scan="page")
    assert [{b.text_source for b in p.blocks} for p in doc.pages] == [{"pdf_text"}, {"vlm"}, {"vlm"}]
    assert doc.meta["calls"] == {"structure": 1, "ocr": 0, "vlm": 2} and len(seen) == 2
    assert doc.meta["vlm"]["enabled"] and doc.meta["vlm_fallbacks"] == []
    scanned = [b for b in doc.pages[1].blocks]
    assert scanned and all(b.bbox == [0.0, 0.0, 595.0, 842.0] for b in scanned)
    validate_doc(doc)
    for name in ("image.png", "image.bmp", "image.webp", "multi.tiff"):
        d = default(str(FIX / name), backend=FakeBackend(), vlm=cfg, vlm_scan="page")
        assert srcs(d) == {"vlm"} and d.meta["calls"]["vlm"] == len(d.pages) and d.meta["calls"]["ocr"] == 0, name
        validate_doc(d)


def test_digital_only_makes_no_vlm_calls(stub):
    cfg, seen = stub
    for name in ("digital.pdf", "contract_ko.pdf"):
        doc = default(str(FIX / name), backend="heuristic", vlm=cfg)
        assert doc.meta["calls"]["vlm"] == 0 and srcs(doc) == {"pdf_text"}
    assert seen == []


def test_figure_blocks_get_vlm_description(stub):
    cfg, seen = stub
    doc = default(str(FIX / "figure.pdf"), backend="heuristic", vlm=cfg)
    figs = [b for _, b in doc.iter_blocks() if b.type == "figure"]
    assert len(figs) == 1 and figs[0].text_source == "vlm" and figs[0].text.startswith("Figure: ")
    assert doc.meta["calls"]["vlm"] == 1 and not [b for _, b in doc.iter_blocks() if b.type == "table"]
    assert any("Quarterly" in (b.text or "") and b.text_source == "pdf_text" for _, b in doc.iter_blocks())
    validate_doc(doc)
    off = default(str(FIX / "figure.pdf"), backend="heuristic", vlm=False)
    assert off.meta["calls"]["vlm"] == 0 and off.meta["vlm"]["enabled"] is False
    assert not [b for _, b in off.iter_blocks() if b.type == "figure" and b.text_source == "vlm"]


def test_pptx_picture_described(stub):
    cfg, _ = stub
    doc = default(str(FIX / "slides.pptx"), vlm=cfg)
    (_, fig), = [(p, b) for p, b in doc.iter_blocks() if b.type == "figure"]
    assert fig.text.startswith("Figure: 240x160") and fig.text_source == "vlm"
    assert doc.meta["calls"] == {"structure": 0, "ocr": 0, "vlm": 1}
    validate_doc(doc)


class WeakTableBackend(FakeBackend):
    """Reports a table region without any cells -> below threshold -> VLM table prompt."""

    def analyze(self, page_image, words):
        s = super().analyze(page_image, words)
        if words:
            s.regions.append(Region(type="table", bbox=[100, 100, 900, 400]))
        return s


def test_weak_table_redone_by_vlm(stub):
    cfg, _ = stub
    doc = default(str(FIX / "digital.pdf"), backend=WeakTableBackend(), vlm=cfg)
    tb = [b for _, b in doc.iter_blocks() if b.type == "table"]
    assert tb and all(b.text_source == "vlm" and b.cells == [["col1", "col2"], ["a", "1"], ["b", "2"]] for b in tb)
    assert doc.meta["calls"]["vlm"] == len(tb)
    validate_doc(doc)
    plain = default(str(FIX / "table.pdf"), backend="heuristic", vlm=cfg)  # good structure -> untouched
    assert plain.meta["calls"]["vlm"] == 0


def test_vlm_http_failure_falls_back_to_ocr():
    def boom(request):
        return httpx.Response(500, text="down")

    client = httpx.Client(transport=httpx.MockTransport(boom))
    fb = FakeBackend(ocr_words_for(["스캔 문서", "OCR text"]))
    doc = default(str(FIX / "mixed.pdf"), backend=fb, vlm={"url": "http://x/v1", "model": "m", "client": client},
                  vlm_scan="page")
    assert [{b.text_source for b in p.blocks} for p in doc.pages] == [{"pdf_text"}, {"ocr"}, {"ocr"}]
    assert doc.meta["errors"] == [] and doc.meta["calls"]["vlm"] == 2 and doc.meta["calls"]["ocr"] == 2
    fbs = doc.meta["vlm_fallbacks"]
    assert [(f["page"], f["task"]) for f in fbs] == [(2, "page"), (3, "page")] and "500" in fbs[0]["error"]
    validate_doc(doc)


def test_vlm_failure_and_ocr_failure_is_page_error():
    def bad(png):
        raise TimeoutError("slow")

    doc = default(str(FIX / "image.png"), backend=FakeBackend(fail_ocr_on=1), vlm=bad)
    assert doc.meta["vlm_fallbacks"][0]["error"].startswith("vlm: TimeoutError")
    assert [e["page"] for e in doc.meta["errors"]] == [1] and not doc.pages[0].blocks


def test_custom_callable_and_off_values():
    def vlm(png):
        return "# 제목\n\n본문 문단\n\n- a\n- b\n\n| x | y |\n|---|---|\n| 1 | 2 |\n"

    doc = default(str(FIX / "scanned.pdf"), backend=FakeBackend(), vlm=vlm)
    assert [b.type for _, b in doc.iter_blocks()] == ["heading", "paragraph", "list", "table"]
    assert srcs(doc) == {"vlm"} and doc.meta["calls"]["vlm"] == 1
    validate_doc(doc)
    for off in (False, None, "off"):
        fb = FakeBackend(ocr_words_for(["abc"]))
        d = default(str(FIX / "scanned.pdf"), backend=fb, vlm=off)
        assert d.meta["calls"]["vlm"] == 0 and srcs(d) == {"ocr"} and d.meta["vlm"]["enabled"] is False


def test_stub_server_prompts():
    c = TestClient(vlm_stub.create_app())
    assert c.get("/v1/models").json()["data"]
    from folio.parse.vlm import PROMPTS, make_caller

    call = make_caller({"url": "http://testserver/v1", "model": "m", "client": c})
    png = (FIX / "image.png").read_bytes()
    assert call(png, task="figure").startswith("Figure: ")
    assert call(png, task="table").startswith("| col1")
    assert "Scanned" in call(png, task="page") or "스캔" in call(png, task="page") or call(png).strip()
    assert call(None, task="table_header", text="0: a | b\n1: 1 | 2") == '{"header_rows": 1}'
    assert set(PROMPTS) == {"page", "figure", "table", "table_header"}


def _low_conf(words, conf):
    return [w.model_copy(update={"conf": conf}) for w in words]


class OcrTableBackend(FakeBackend):
    """OCR pages: preset words plus one table region with a 1x1 cell over them."""

    def analyze(self, page_image, words):
        s = super().analyze(page_image, words)
        if words is None:
            s.regions.append(Region(type="table", bbox=[90, 90, 600, 200]))
            s.tables.append(Table(bbox=[90, 90, 600, 200], cells=[TableCell(bbox=[90, 90, 600, 200], row=0, col=0)]))
        return s


def test_ocr_first_keeps_ocr_bboxes_and_redoes_tables(stub):
    cfg, seen = stub
    fb = OcrTableBackend(ocr_words_for(["스캔 문서", "OCR text"]))
    doc = default(str(FIX / "image.png"), backend=fb, vlm=cfg, vlm_scan="ocr_first")
    blocks = doc.pages[0].blocks
    assert {b.text_source for b in blocks} == {"vlm"} and [b.type for b in blocks] == ["table"]
    assert blocks[0].bbox != [0.0, 0.0, doc.pages[0].width, doc.pages[0].height]
    assert doc.meta["calls"]["ocr"] == 1 and seen == ["/v1/chat/completions"]  # one table request, no page request
    validate_doc(doc)
    plain = default(str(FIX / "image.png"), backend=FakeBackend(ocr_words_for(["스캔 문서", "OCR text"])), vlm=cfg,
                    vlm_scan="ocr_first")
    assert srcs(plain) == {"ocr"} and plain.meta["calls"]["vlm"] == 0


def test_ocr_first_low_confidence_goes_to_vlm_page(stub):
    cfg, _ = stub
    fb = FakeBackend(_low_conf(ocr_words_for(["흐린 스캔"]), 0.4))
    doc = default(str(FIX / "image.png"), backend=fb, vlm=cfg, vlm_scan="ocr_first")
    assert srcs(doc) == {"vlm"} and doc.meta["calls"] == {"structure": 0, "ocr": 1, "vlm": 1}
    off = default(str(FIX / "image.png"), backend=FakeBackend(_low_conf(ocr_words_for(["흐린 스캔"]), 0.4)),
                  vlm=False, vlm_scan="ocr_first")  # no VLM: plain OCR, as before
    assert srcs(off) == {"ocr"}
    with pytest.raises(ValueError):
        default(str(FIX / "image.png"), backend=FakeBackend(), vlm=cfg, vlm_scan="nope")


class OcrFigureBackend(FakeBackend):
    """OCR pages: preset words inside one figure region (e.g. a table the layout model called a figure)."""

    def analyze(self, page_image, words):
        s = super().analyze(page_image, words)
        if words is None:
            s.regions = [Region(type="figure", bbox=[90, 90, 600, 200])]
        return s


def test_figure_description_keeps_text_read_in_region(stub):
    cfg, _ = stub
    doc = default(str(FIX / "image.png"), backend=OcrFigureBackend(ocr_words_for(["매출 100억", "영업이익 5억"])),
                  vlm=cfg, vlm_scan="ocr_first")
    (fig,) = [b for _, b in doc.iter_blocks() if b.type == "figure"]
    assert fig.text.startswith("매출 100억") and "Figure: " in fig.text and fig.text_source == "ocr"


def test_vlm_html_table_keeps_merged_cells():
    def vlm(png, task="page"):
        return "<table><tr><td colspan=2>합계</td></tr><tr><td>a</td><td>1</td></tr></table>" if task == "table" else ""

    doc = default(str(FIX / "digital.pdf"), backend=WeakTableBackend(), vlm=vlm)
    tb = [b for _, b in doc.iter_blocks() if b.type == "table"]
    assert tb and all(b.text_source == "vlm" and b.cells == [["합계", ""], ["a", "1"]] for b in tb)
    assert all([(s.row, s.col, s.rowspan, s.colspan) for s in b.spans] == [(0, 0, 1, 2)] for b in tb)
    validate_doc(doc)


def _xlsx(build) -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet_tables(tmp_path, data, **kw):
    p = tmp_path / "t.xlsx"
    p.write_bytes(data)
    doc = default(str(p), **kw)
    validate_doc(doc)
    return doc, [b for _, b in doc.iter_blocks() if b.type == "table"]


def test_xlsx_header_rows_native_sources_skip_vlm(stub, tmp_path):
    from openpyxl.worksheet.table import Table as XTable

    def build(wb):
        ws = wb.active
        for r in [["분기 실적", ""], ["지역", "매출"], ["서울", 10], ["부산", 7]]:
            ws.append(r)
        ws.add_table(XTable(displayName="T1", ref="A2:B4"))  # title row above the Excel table counts as header
        ws2 = wb.create_sheet("p")
        for r in [["구분", "값"], ["단위", "원"], ["a", 1], ["b", 2]]:
            ws2.append(r)
        ws2.print_title_rows = "1:2"

    cfg, seen = stub
    doc, (t1, t2) = _sheet_tables(tmp_path, _xlsx(build), vlm=cfg)
    assert (t1.header_rows, t2.header_rows) == (2, 2)
    assert doc.meta["calls"]["vlm"] == 0 and seen == []


def test_xlsx_header_rows_asked_per_region(stub, tmp_path):
    def build(wb):
        ws = wb.active
        for r in [["제품", "수량"], ["사과", 1], ["배", 2], [], ["합계 메모"]]:
            ws.append(r)

    cfg, seen = stub
    doc, (t, note) = _sheet_tables(tmp_path, _xlsx(build), vlm=cfg)
    assert t.header_rows == 1 and note.header_rows is None  # 1x1 region is not asked
    assert doc.meta["calls"]["vlm"] == 1 and seen == ["/v1/chat/completions"]
    _, (t, _) = _sheet_tables(tmp_path, _xlsx(build), vlm=cfg, xlsx_header_llm=False)
    assert t.header_rows is None


def test_xlsx_header_reply_parsing_and_image_only_caller(tmp_path):
    def build(wb):
        for r in [["a", "b"], ["c", "d"], ["1", "2"]]:
            wb.active.append(r)

    data = _xlsx(build)
    replies = {"<think>header_rows: 1?</think>```json\n{\"header_rows\": 2}\n```": 2, '{"header_rows": 3}': None,
               "모르겠음": None}
    for reply, want in replies.items():
        doc, (t,) = _sheet_tables(tmp_path, data, vlm=lambda png, task="page", text=None, r=reply: r)
        assert t.header_rows == want
        assert len(doc.meta["vlm_fallbacks"]) == (reply == "모르겠음")
    doc, (t,) = _sheet_tables(tmp_path, data, vlm=lambda png: "x")  # image-only caller: skipped silently
    assert t.header_rows is None and doc.meta["calls"]["vlm"] == 0 and doc.meta["vlm_fallbacks"] == []
